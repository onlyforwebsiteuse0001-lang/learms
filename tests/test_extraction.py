"""Real document parsing and explicit OCR-failure tests."""
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from reportlab.pdfgen import canvas

from backend.app.core.config import Settings
from backend.app.services.extraction_service import (
    ExtractionError,
    extract_from_docx,
    extract_from_image,
    extract_from_pdf,
)
from backend.app.services.ocr_service import OCRResult
from backend.app.services.text_cleaner import clean_pages, detect_language


@pytest.fixture
def digital_pdf(tmp_path: Path) -> Path:
    """Create a small valid PDF containing an actual digital text stream."""
    path = tmp_path / "digital.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 760, "Photosynthesis converts light energy into chemical energy.")
    pdf.drawString(72, 740, "Chlorophyll absorbs light in plant cells.")
    pdf.save()
    return path


def test_extract_pdf_digital(digital_pdf: Path) -> None:
    """Digital PDF text must be extracted by a local PDF parser."""
    result = extract_from_pdf(digital_pdf, Settings(_env_file=None, gemini_api_key=None))
    assert "Photosynthesis" in result.pages[0].content
    assert result.pages[0].method in {"pdfplumber", "pypdf2"}


def test_extract_pdf_scanned_uses_ocr(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """A page without embedded text must take the Tesseract route."""
    path = tmp_path / "scan.pdf"
    Image.new("RGB", (800, 500), "white").save(path, "PDF")
    monkeypatch.setattr("backend.app.services.extraction_service.convert_from_path", lambda *args, **kwargs: [Image.new("RGB", (800, 500), "white")])
    monkeypatch.setattr("backend.app.services.extraction_service.run_tesseract", lambda image: OCRResult("Scanned biology notes from a real OCR result", 91.2))
    result = extract_from_pdf(path, Settings(_env_file=None, gemini_api_key=None))
    assert result.pages[0].method == "tesseract"
    assert result.pages[0].confidence == 91.2


def test_extract_docx(tmp_path: Path) -> None:
    """DOCX paragraphs and tables must be extracted without an AI provider."""
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_heading("Cell Biology", level=1)
    document.add_paragraph("Mitochondria produce ATP through cellular respiration.")
    document.save(path)
    result = extract_from_docx(path)
    assert "Mitochondria" in result.pages[0].content
    assert result.pages[0].method == "docx"


def test_extract_image_uses_tesseract(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """Supported images must route to Tesseract and retain confidence."""
    path = tmp_path / "notes.png"
    Image.new("RGB", (600, 300), "white").save(path)
    monkeypatch.setattr("backend.app.services.extraction_service.run_tesseract", lambda image: OCRResult("Newton second law force equals mass times acceleration", 88.4))
    result = extract_from_image(path, Settings(_env_file=None, gemini_api_key=None))
    assert result.pages[0].method == "tesseract"
    assert result.pages[0].confidence == 88.4


def test_extract_failure_no_fake_output(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """Local and AI failure must raise an error rather than manufacture text."""
    path = tmp_path / "blank.png"
    Image.new("RGB", (50, 50), "white").save(path)
    monkeypatch.setattr("backend.app.services.extraction_service.run_tesseract", lambda image: OCRResult("", None))
    with pytest.raises(ExtractionError, match="AI vision fallback unavailable"):
        extract_from_image(path, Settings(_env_file=None, gemini_api_key=None))


def test_text_cleaning_removes_repeated_headers_footers_and_numbers() -> None:
    """Repeated boundaries and standalone page labels should be removed."""
    pages = [
        "Biology 101\nCell structure and membrane transport.\nPage 1",
        "Biology 101\nMitochondria and energy production.\nPage 2",
        "Biology 101\nDNA replication and inheritance.\nPage 3",
    ]
    cleaned = clean_pages(pages)
    assert all("Biology 101" not in page and "Page" not in page for page in cleaned)
    assert "Cell structure" in cleaned[0]


@pytest.mark.parametrize(("text", "expected"), [
    ("Cell biology studies living organisms", "en"),
    ("خلیہ زندگی کی بنیادی اکائی ہے", "ur"),
    ("Cell کو خلیہ کہتے ہیں", "mixed"),
    ("12 + 4", "unknown"),
])
def test_language_detection(text: str, expected: str) -> None:
    """English, Urdu, mixed script, and unknown content are distinguished."""
    assert detect_language(text) == expected
