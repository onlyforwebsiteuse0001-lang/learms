"""Failure, fallback and contract tests for document and concept extraction."""
from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation

from ai_services.concept_router import (
    _gemini,
    _groq,
    _json_object,
    _openrouter,
    _validate,
    extract_concepts,
)
from backend.app.core.config import Settings
from backend.app.services.extraction_service import (
    ExtractionError,
    _meaningful,
    extract_from_docx,
    extract_from_image,
    extract_from_pdf,
    extract_from_pptx,
    extract_text,
)


@pytest.mark.parametrize(
    ("text", "expected"),
    [("", False), (" \n\t", False), ("123456789", False), ("12345 67890", True), ("خلیہ بنیادی اکائی", True), ("a" * 100_000, True)],
)
def test_meaningful_text_boundary(text: str, expected: bool) -> None:
    assert _meaningful(text) is expected


@pytest.mark.parametrize("raw", ["", "plain prose", "[]", "```json\n[]\n```", "{bad json}"])
def test_provider_json_parser_rejects_invalid_contract(raw: str) -> None:
    with pytest.raises((ValueError, json.JSONDecodeError)):
        _json_object(raw)


def test_provider_json_parser_accepts_fenced_or_prefixed_object() -> None:
    assert _json_object('result: ```json\n{"concepts": []}\n```')["concepts"] == []


def _payload(evidence: str = "Cell Biology", *, duplicate: bool = False, relation_evidence: str = "Genetics requires Cell Biology.") -> str:
    concepts = [
        {"name": "Cell Biology", "description": "Study of cells", "evidence": evidence, "confidence": .8},
        {"name": "Genetics", "description": "Study of inheritance", "evidence": "Genetics", "confidence": .9},
    ]
    if duplicate:
        concepts.append({"name": " cell biology! ", "description": "Duplicate", "evidence": evidence, "confidence": .7})
    return json.dumps({"concepts": concepts, "prerequisites": [{"prerequisite": "Cell Biology", "concept": "Genetics", "evidence": relation_evidence, "confidence": .85}]})


def test_provider_validation_keeps_only_source_grounded_unique_data() -> None:
    text = "Cell Biology\nGenetics\nGenetics requires Cell Biology."
    result = _validate(_payload(duplicate=True), text, "mock_structured")
    assert len(result.concepts) == 2
    assert len(result.prerequisites) == 1
    assert all(item.evidence in text for item in [*result.concepts, *result.prerequisites])


@pytest.mark.parametrize(
    "mutation",
    [
        {"concepts": []},
        {"concepts": [{"name": "X", "description": "Y", "evidence": "invented", "confidence": .5}]},
        {"concepts": [{"name": "X", "description": "Y", "evidence": "X", "confidence": 2}]},
        {"not_concepts": []},
    ],
)
def test_provider_validation_rejects_empty_partial_or_invalid_payload(mutation: dict) -> None:
    with pytest.raises((ValueError, KeyError)):
        _validate(json.dumps(mutation), "X", "mock")


def test_invalid_relation_is_dropped_without_losing_concepts() -> None:
    text = "Cell Biology\nGenetics"
    result = _validate(_payload(relation_evidence="not in source"), text, "mock")
    assert len(result.concepts) == 2 and result.prerequisites == []


def test_provider_order_gemini_failure_then_groq_success(monkeypatch: pytest.MonkeyPatch) -> None:
    text = "Cell Biology\nGenetics\nGenetics requires Cell Biology."
    monkeypatch.setattr("ai_services.concept_router._gemini", lambda *_: (_ for _ in ()).throw(TimeoutError("timeout")))
    monkeypatch.setattr("ai_services.concept_router._groq", lambda *_: _payload())
    result = extract_concepts(text, Settings(_env_file=None, gemini_api_key="g", groq_api_key="q"))
    assert result.method == "groq_structured"


def test_provider_order_reaches_openrouter(monkeypatch: pytest.MonkeyPatch) -> None:
    text = "Cell Biology\nGenetics\nGenetics requires Cell Biology."
    monkeypatch.setattr("ai_services.concept_router._gemini", lambda *_: "invalid")
    monkeypatch.setattr("ai_services.concept_router._groq", lambda *_: "also invalid")
    monkeypatch.setattr("ai_services.concept_router._openrouter", lambda *_: _payload())
    settings = Settings(_env_file=None, gemini_api_key="g", groq_api_key="q", openrouter_api_key="o")
    assert extract_concepts(text, settings).method == "openrouter_structured"


def test_all_provider_failures_are_recorded_before_deterministic_fallback(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("ai_services.concept_router._gemini", lambda *_: (_ for _ in ()).throw(RuntimeError("down")))
    monkeypatch.setattr("ai_services.concept_router._groq", lambda *_: "bad")
    monkeypatch.setattr("ai_services.concept_router._openrouter", lambda *_: "{}")
    text = "Cell Biology\nThis is a sufficiently long description about cells and organisms."
    settings = Settings(_env_file=None, gemini_api_key="g", groq_api_key="q", openrouter_api_key="o")
    result = extract_concepts(text, settings)
    assert result.method == "deterministic_source_only"
    assert len(result.provider_errors) == 3


def test_gemini_adapter_contract(monkeypatch: pytest.MonkeyPatch) -> None:
    captured = {}
    response = SimpleNamespace(text="result")
    models = SimpleNamespace(generate_content=lambda **kwargs: captured.update(kwargs) or response)
    monkeypatch.setattr("google.genai.Client", lambda api_key: SimpleNamespace(models=models))
    output = _gemini("source", Settings(_env_file=None, gemini_api_key="secret"))
    assert output == "result" and "SOURCE" in captured["contents"]


def test_groq_adapter_contract(monkeypatch: pytest.MonkeyPatch) -> None:
    captured = {}
    completions = SimpleNamespace(create=lambda **kwargs: captured.update(kwargs) or SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="result"))]))
    monkeypatch.setattr("groq.Groq", lambda api_key: SimpleNamespace(chat=SimpleNamespace(completions=completions)))
    assert _groq("source", Settings(_env_file=None, groq_api_key="secret")) == "result"
    assert captured["response_format"] == {"type": "json_object"}


def test_openrouter_adapter_contract(monkeypatch: pytest.MonkeyPatch) -> None:
    captured = {}
    response = SimpleNamespace(raise_for_status=lambda: None, json=lambda: {"choices": [{"message": {"content": "result"}}]})
    monkeypatch.setattr("ai_services.concept_router.httpx.post", lambda *args, **kwargs: captured.update(kwargs) or response)
    assert _openrouter("source", Settings(_env_file=None, openrouter_api_key="secret")) == "result"
    assert captured["timeout"] == 45


def test_corrupt_docx_fails_without_fake_text(tmp_path: Path) -> None:
    path = tmp_path / "bad.docx"; path.write_bytes(b"not a zip")
    with pytest.raises(ExtractionError, match="DOCX extraction failed"):
        extract_from_docx(path)


def test_empty_pptx_fails_without_fake_text(tmp_path: Path) -> None:
    path = tmp_path / "empty.pptx"; Presentation().save(path)
    with pytest.raises(ExtractionError, match="no extractable text"):
        extract_from_pptx(path)


def test_pptx_extracts_slide_text_and_table(tmp_path: Path) -> None:
    path = tmp_path / "lesson.pptx"
    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Cell Biology"
    table = slide.shapes.add_table(1, 1, 10, 10, 100, 100).table; table.cell(0, 0).text = "Mitochondria produce cellular energy"
    presentation.save(path)
    result = extract_from_pptx(path)
    assert result.page_count == 1 and "Mitochondria" in result.pages[0].content


@pytest.mark.parametrize("suffix", [".txt", ".exe", ".zip", ""])
def test_extract_text_rejects_unsupported_suffix(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / f"file{suffix}"; path.write_text("content")
    with pytest.raises(ExtractionError, match="No extractor"):
        extract_text(path, "application/octet-stream", Settings(_env_file=None))


def test_image_local_exception_and_unavailable_ai_are_reported(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "image.png"; Image.new("RGB", (10, 10)).save(path)
    monkeypatch.setattr("backend.app.services.extraction_service.run_tesseract", lambda *_: (_ for _ in ()).throw(RuntimeError("engine down")))
    with pytest.raises(ExtractionError, match="Tesseract failed: engine down"):
        extract_from_image(path, Settings(_env_file=None, gemini_api_key=None))


def test_pdf_all_local_extractors_and_ai_failure_are_explicit(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"; path.write_bytes(b"not pdf")
    monkeypatch.setattr("backend.app.services.extraction_service._pdfplumber_pages", lambda *_: (_ for _ in ()).throw(RuntimeError("plumber down")))
    monkeypatch.setattr("backend.app.services.extraction_service._pypdf_pages", lambda *_: (_ for _ in ()).throw(RuntimeError("reader down")))
    monkeypatch.setattr("backend.app.services.extraction_service.convert_from_path", lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("renderer down")))
    with pytest.raises(ExtractionError, match="plumber down"):
        extract_from_pdf(path, Settings(_env_file=None, gemini_api_key=None))


def test_pdf_uses_pypdf_when_plumber_has_no_meaningful_text(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    path = tmp_path / "mock.pdf"; path.write_bytes(b"mock")
    monkeypatch.setattr("backend.app.services.extraction_service._pdfplumber_pages", lambda *_: [""])
    monkeypatch.setattr("backend.app.services.extraction_service._pypdf_pages", lambda *_: ["A meaningful source sentence from pypdf fallback."])
    result = extract_from_pdf(path, Settings(_env_file=None))
    assert result.pages[0].method == "pypdf2"
